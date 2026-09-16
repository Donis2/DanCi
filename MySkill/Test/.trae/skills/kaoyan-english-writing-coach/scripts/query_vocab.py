#!/usr/bin/env python3
# query_vocab.py - 考研英语大纲词汇查询工具
#
# 用途：辅助批改学生短文时，查询某个词是否在考研大纲词汇表中，
#       以及查找同义/相关的高级替换词（按词频排序）。
#
# 数据文件：考研英语大纲词汇.xlsx
#   列：序号、词频、单词、释义、其他拼写、分类、子分类
#
# 用法：
#   python query_vocab.py --xlsx <path> <命令> [参数]
#
# 命令：
#   lookup <word>           查询单词是否在大纲中，返回词频/释义/分类
#   lookup-bulk w1,w2,w3    批量查询多个单词
#   search <keyword>        模糊搜索（按单词或释义包含关键词）
#   freq-top <category>     列出某分类下词频最高的 N 个词
#   freq-range <min> <max>  列出词频在 [min, max] 范围内的词
#   stats                   输出词汇表统计信息
#
# 示例：
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx lookup appropriate
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx lookup-bulk proper,significant,important
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx search proper
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx freq-top 动作行为 --limit 20
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx freq-range 100 200 --limit 30
#   python query_vocab.py --xlsx 考研英语大纲词汇.xlsx stats

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("[ERROR] 需要 openpyxl 库，请运行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def load_workbook(xlsx_path: str):
    """加载 xlsx，返回 (header, rows) 其中 rows 是 dict 列表"""
    p = Path(xlsx_path)
    if not p.exists():
        print(f"[ERROR] 文件不存在: {xlsx_path}", file=sys.stderr)
        sys.exit(1)
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    header = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(header)]
    data = []
    for row in rows_iter:
        if row is None or all(c is None for c in row):
            continue
        data.append(dict(zip(header, row)))
    wb.close()
    return header, data


def cmd_lookup(args, data):
    """查询单词（精确匹配，忽略大小写）"""
    target = args.word.strip().lower()
    for r in data:
        w = str(r.get("单词", "")).strip()
        if w.lower() == target:
            print(f"✓ 找到: {w}")
            print(f"  序号: {r.get('序号')}")
            print(f"  词频: {r.get('词频')}")
            print(f"  释义: {r.get('释义')}")
            print(f"  其他拼写: {r.get('其他拼写') or '-'}")
            print(f"  分类: {r.get('分类')} / {r.get('子分类')}")
            return
    print(f"✗ 未找到: {args.word}（不在考研大纲词汇中）")


def cmd_lookup_bulk(args, data):
    """批量查询多个单词，用逗号分隔"""
    words = [w.strip() for w in args.words.split(",") if w.strip()]
    if not words:
        print("[ERROR] 请用逗号分隔多个单词")
        return
    print(f"批量查询 {len(words)} 个单词：\n")
    word_set = {w.lower() for w in words}
    found = {}
    for r in data:
        w = str(r.get("单词", "")).strip()
        if w.lower() in word_set:
            found[w.lower()] = r
    for w in words:
        r = found.get(w.lower())
        if r:
            print(f"✓ {w:<20} 词频 {r.get('词频'):<6} {r.get('释义', '')[:50]}")
        else:
            print(f"✗ {w:<20} 不在大纲中")


def cmd_search(args, data):
    """模糊搜索：单词或释义包含关键词"""
    kw = args.keyword.strip().lower()
    if not kw:
        print("[ERROR] 请提供关键词")
        return
    matches = []
    for r in data:
        w = str(r.get("单词", "")).strip()
        d = str(r.get("释义", "") or "").strip()
        if kw in w.lower() or kw in d.lower():
            matches.append(r)
    if not matches:
        print(f"未找到包含 '{args.keyword}' 的词条")
        return
    print(f"找到 {len(matches)} 条匹配 '{args.keyword}':\n")
    for r in matches[:args.limit]:
        print(f"  {r.get('单词'):<20} 词频 {r.get('词频'):<6} {r.get('释义', '')[:60]}")


def cmd_freq_top(args, data):
    """列出某分类下词频最高的 N 个词"""
    cat = args.category.strip()
    matches = [r for r in data if str(r.get("分类", "") or "").strip() == cat]
    if not matches:
        print(f"未找到分类: {cat}")
        print("可用分类:", ", ".join(sorted({str(r.get('分类', '') or '') for r in data})))
        return
    matches.sort(key=lambda r: -(int(r.get("词频") or 0)))
    print(f"分类 '{cat}' 下词频最高的 {min(args.limit, len(matches))} 个词:\n")
    for r in matches[:args.limit]:
        print(f"  {r.get('单词'):<20} 词频 {r.get('词频'):<6} {r.get('释义', '')[:50]}")


def cmd_freq_range(args, data):
    """列出词频在 [min, max] 范围内的词"""
    try:
        lo = int(args.min)
        hi = int(args.max)
    except ValueError:
        print("[ERROR] min 和 max 必须是整数")
        return
    matches = []
    for r in data:
        try:
            f = int(r.get("词频") or 0)
        except (ValueError, TypeError):
            continue
        if lo <= f <= hi:
            matches.append(r)
    matches.sort(key=lambda r: -int(r.get("词频") or 0))
    print(f"词频在 [{lo}, {hi}] 范围内的词（共 {len(matches)} 个，显示前 {args.limit}）:\n")
    for r in matches[:args.limit]:
        print(f"  {r.get('单词'):<20} 词频 {r.get('词频'):<6} {r.get('释义', '')[:50]}")


def cmd_stats(args, data):
    """输出词汇表统计信息"""
    total = len(data)
    by_cat = {}
    for r in data:
        cat = str(r.get("分类", "") or "未分类")
        by_cat[cat] = by_cat.get(cat, 0) + 1
    print(f"考研英语大纲词汇统计")
    print(f"=" * 40)
    print(f"总词数: {total}")
    print(f"\n按分类分布:")
    for cat, cnt in sorted(by_cat.items(), key=lambda x: -x[1]):
        print(f"  {cat:<15} {cnt} 词")


def main():
    parser = argparse.ArgumentParser(description="考研英语大纲词汇查询工具")
    parser.add_argument("--xlsx", required=True, help="考研英语大纲词汇.xlsx 路径")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_lookup = sub.add_parser("lookup", help="查询单词")
    p_lookup.add_argument("word")

    p_bulk = sub.add_parser("lookup-bulk", help="批量查询（逗号分隔）")
    p_bulk.add_argument("words")

    p_search = sub.add_parser("search", help="模糊搜索")
    p_search.add_argument("keyword")
    p_search.add_argument("--limit", type=int, default=20)

    p_top = sub.add_parser("freq-top", help="某分类下词频最高的词")
    p_top.add_argument("category")
    p_top.add_argument("--limit", type=int, default=20)

    p_range = sub.add_parser("freq-range", help="词频范围内的词")
    p_range.add_argument("min")
    p_range.add_argument("max")
    p_range.add_argument("--limit", type=int, default=30)

    sub.add_parser("stats", help="词汇表统计")

    args = parser.parse_args()
    header, data = load_workbook(args.xlsx)

    if args.cmd == "lookup":
        cmd_lookup(args, data)
    elif args.cmd == "lookup-bulk":
        cmd_lookup_bulk(args, data)
    elif args.cmd == "search":
        cmd_search(args, data)
    elif args.cmd == "freq-top":
        cmd_freq_top(args, data)
    elif args.cmd == "freq-range":
        cmd_freq_range(args, data)
    elif args.cmd == "stats":
        cmd_stats(args, data)


if __name__ == "__main__":
    main()
